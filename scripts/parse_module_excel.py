import sys
import json
import os
import openpyxl

def parse_excel(file_path):
    if not os.path.exists(file_path):
        return {"success": False, "error": f"File not found: {file_path}"}

    try:
        wb = openpyxl.load_workbook(file_path, data_only=True)
    except Exception as e:
        return {"success": False, "error": f"Failed to open Excel file: {str(e)}"}

    modules = []
    topics = []
    questions = []

    # 1. Parse Modul Utama
    module_sheet_name = None
    for name in wb.sheetnames:
        if "modul" in name.lower() and "utama" in name.lower():
            module_sheet_name = name
            break
    if not module_sheet_name and len(wb.sheetnames) > 0:
        module_sheet_name = wb.sheetnames[0]

    if module_sheet_name in wb.sheetnames:
        ws = wb[module_sheet_name]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or row[0] is None:
                continue
            # Column mapping:
            # 0: Nomor Modul, 1: Judul Modul, 2: Sub-Judul, 3: Deskripsi, 4: Estimasi Waktu, 5: Banner, 6: Badge Icon
            try:
                mod_num = int(row[0])
            except (ValueError, TypeError):
                continue

            title = str(row[1]).strip() if len(row) > 1 and row[1] is not None else ""
            if not title:
                continue

            subtitle = str(row[2]).strip() if len(row) > 2 and row[2] is not None else None
            desc = str(row[3]).strip() if len(row) > 3 and row[3] is not None else ""
            est_time = str(row[4]).strip() if len(row) > 4 and row[4] is not None else "15 Menit"
            banner = str(row[5]).strip() if len(row) > 5 and row[5] is not None else None
            badge = str(row[6]).strip() if len(row) > 6 and row[6] is not None else None

            modules.append({
                "module_number": mod_num,
                "title": title,
                "subtitle": subtitle,
                "description": desc,
                "estimated_time": est_time,
                "banner_image": banner,
                "badge_icon": badge
            })

    # 2. Parse Submateri (Topik)
    topic_sheet_name = None
    for name in wb.sheetnames:
        if "topik" in name.lower() or "submateri" in name.lower():
            topic_sheet_name = name
            break
    if not topic_sheet_name and len(wb.sheetnames) > 1:
        topic_sheet_name = wb.sheetnames[1]

    if topic_sheet_name in wb.sheetnames:
        ws = wb[topic_sheet_name]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or row[0] is None:
                continue
            # 0: Modul Target, 1: Kode Topik, 2: Urutan, 3: Judul, 4: Video ID, 5: Konten HTML
            try:
                mod_target = int(row[0])
            except (ValueError, TypeError):
                continue

            topic_code = str(row[1]).strip() if len(row) > 1 and row[1] is not None else f"{mod_target}.1"
            try:
                order_index = int(row[2]) if len(row) > 2 and row[2] is not None else 1
            except (ValueError, TypeError):
                order_index = 1

            title = str(row[3]).strip() if len(row) > 3 and row[3] is not None else ""
            if not title:
                continue

            video_id = str(row[4]).strip() if len(row) > 4 and row[4] is not None else None
            if video_id and ("http" in video_id or "youtube.com" in video_id or "youtu.be" in video_id):
                # Ekstrak video id jika user memasukkan full url
                if "v=" in video_id:
                    video_id = video_id.split("v=")[1].split("&")[0]
                elif "youtu.be/" in video_id:
                    video_id = video_id.split("youtu.be/")[1].split("?")[0]

            content_html = str(row[5]).strip() if len(row) > 5 and row[5] is not None else ""

            topics.append({
                "module_target": mod_target,
                "topic_code": topic_code,
                "order_index": order_index,
                "title": title,
                "youtube_video_id": video_id,
                "content_html": content_html
            })

    # 3. Parse Bank Soal Kuesioner
    quiz_sheet_name = None
    for name in wb.sheetnames:
        if "soal" in name.lower() or "kuesioner" in name.lower():
            quiz_sheet_name = name
            break
    if not quiz_sheet_name and len(wb.sheetnames) > 2:
        quiz_sheet_name = wb.sheetnames[2]

    if quiz_sheet_name in wb.sheetnames:
        ws = wb[quiz_sheet_name]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or row[0] is None:
                continue
            # 0: Modul Target, 1: Nomor Soal, 2: Teks Pertanyaan, 3: Opsi A, 4: Opsi B, 5: Opsi C, 6: Opsi D, 7: Kunci, 8: Penjelasan
            try:
                mod_target = int(row[0])
            except (ValueError, TypeError):
                continue

            q_text = str(row[2]).strip() if len(row) > 2 and row[2] is not None else ""
            if not q_text:
                continue

            opt_a = str(row[3]).strip() if len(row) > 3 and row[3] is not None else ""
            opt_b = str(row[4]).strip() if len(row) > 4 and row[4] is not None else ""
            opt_c = str(row[5]).strip() if len(row) > 5 and row[5] is not None else ""
            opt_d = str(row[6]).strip() if len(row) > 6 and row[6] is not None else ""

            correct_ans = str(row[7]).strip().upper() if len(row) > 7 and row[7] is not None else "A"
            if correct_ans not in ["A", "B", "C", "D"]:
                correct_ans = "A"

            explanation = str(row[8]).strip() if len(row) > 8 and row[8] is not None else None

            questions.append({
                "module_target": mod_target,
                "question_text": q_text,
                "options": {
                    "A": opt_a,
                    "B": opt_b,
                    "C": opt_c,
                    "D": opt_d
                },
                "correct_answer": correct_ans,
                "explanation": explanation
            })

    return {
        "success": True,
        "modules": modules,
        "topics": topics,
        "questions": questions
    }

if __name__ == "__main__":
    if len(sys.argv) < 2:
        print(json.dumps({"success": False, "error": "Missing file_path argument"}))
        sys.exit(1)

    file_path = sys.argv[1]
    result = parse_excel(file_path)
    print(json.dumps(result, ensure_ascii=False))
