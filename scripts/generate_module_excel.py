import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def build_excel():
    output_dir = os.path.join(os.path.dirname(__file__), '..', 'public', 'templates')
    os.makedirs(output_dir, exist_ok=True)
    output_file = os.path.join(output_dir, 'format_modul_sikera.xlsx')

    wb = openpyxl.Workbook()
    
    # Color palette based on DESIGN.md
    # Primary Emerald Teal: #0F766E
    # Secondary Accent: #0D9488
    # Dark Header: #1E293B
    # Soft background: #F8FAFC
    
    header_fill_teal = PatternFill(start_color="0F766E", end_color="0F766E", fill_type="solid")
    header_fill_slate = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    sub_fill = PatternFill(start_color="CCFBF1", end_color="CCFBF1", fill_type="solid")
    zebra_fill = PatternFill(start_color="F0FDFA", end_color="F0FDFA", fill_type="solid")
    
    font_header = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    font_title = Font(name="Arial", size=14, bold=True, color="0F766E")
    font_bold = Font(name="Arial", size=10, bold=True, color="000000")
    font_normal = Font(name="Arial", size=10, color="334155")
    font_italic = Font(name="Arial", size=9, italic=True, color="64748B")
    
    thin_border_side = Side(border_style="thin", color="CBD5E1")
    border_cell = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    align_top_left = Alignment(horizontal="left", vertical="top", wrap_text=True)

    # =========================================================================
    # SHEET 1: Modul Utama
    # =========================================================================
    ws1 = wb.active
    ws1.title = "1. Modul Utama"
    ws1.views.sheetView[0].showGridLines = True
    
    # Title Block
    ws1["A1"] = "FORMAT DATA MASTER MODUL EDUKASI SIKERA"
    ws1["A1"].font = font_title
    ws1["A2"] = "Panduan: Gunakan lembar kerja ini untuk menyusun informasi dasar setiap modul edukasi kesehatan reproduksi."
    ws1["A2"].font = font_italic
    
    headers_modul = [
        "Nomor Modul*", "Judul Modul*", "Subjudul / Ringkasan", 
        "Deskripsi Lengkap*", "Estimasi Waktu*", "Path Banner Gambar (Opsional)", "Icon Badge (Opsional)"
    ]
    
    for col_num, h in enumerate(headers_modul, 1):
        cell = ws1.cell(row=4, column=col_num, value=h)
        cell.font = font_header
        cell.fill = header_fill_teal
        cell.alignment = align_center
        cell.border = border_cell
        ws1.row_dimensions[4].height = 28
        
    sample_modules = [
        (1, "Anatomi, Fisiologi, & Higienitas Reproduksi", 
         "Pemahaman dasar struktur biologis, proses pubertas, siklus menstruasi, dan kebersihan diri.",
         "Membahas sistem reproduksi pria dan wanita secara ilmiah tanpa vulgaritas, variasi warna darah haid yang normal vs abnormal, serta penanganan mandiri nyeri haid (dismenore).",
         "15 Menit", "/images/banners/banner-module-1.svg", "anatomy"),
        (2, "Batasan Pergaulan, Pacaran Sehat, & Dinamika Kampus",
         "Membangun relasi yang saling menghargai, komunikasi asertif, dan pencegahan kekerasan relasional.",
         "Membekali mahasiswa baru dengan pemahaman batasan fisik/emosional (boundaries), analisis risiko kos-kosan (living together), serta keterampilan berkata TIDAK secara tegas.",
         "12 Menit", "/images/banners/banner-module-2.svg", "boundaries"),
        (3, "Risiko Medis Seks Bebas & Perlindungan Diri",
         "Edukasi pencegahan IMS, HIV/AIDS, bahaya aborsi ilegal, dan perlindungan keamanan digital.",
         "Tinjauan ilmiah mengenai transmisi patogen seksual, konsekuensi medis aborsi tanpa pengawasan tenaga kesehatan, serta mitigasi Kekerasan Berbasis Gender Online (KBGO).",
         "15 Menit", "/images/banners/banner-module-3.svg", "medical_risk"),
        (4, "Mitigasi Pernikahan Dini & Kesiapan Berkeluarga",
         "Empat pilar kesiapan menikah BKKBN, risiko kehamilan remaja, dan pencegahan stunting.",
         "Mempersiapkan masa depan generasi muda melalui pematangan usia perkawinan (21 tahun wanita, 25 tahun pria), mitigasi risiko panggul sempit (CPD), serta pemenuhan gizi remaja.",
         "10 Menit", "/images/banners/banner-module-4.svg", "family_planning"),
        (5, "[Contoh Modul Baru: Gizi Remaja & Kebugaran]",
         "Nutrisi seimbang untuk kestabilan hormonal reproduksi dan pencegahan anemia.",
         "Membahas pola makan bergizi, kebutuhan zat besi harian, dan olahraga teratur untuk mendukung sistem metabolisme reproduksi.",
         "15 Menit", "/images/banners/banner-module-1.svg", "nutrition")
    ]
    
    for row_idx, mod in enumerate(sample_modules, 5):
        ws1.row_dimensions[row_idx].height = 45
        is_even = (row_idx % 2 == 0)
        for col_idx, val in enumerate(mod, 1):
            cell = ws1.cell(row=row_idx, column=col_idx, value=val)
            cell.font = font_normal
            cell.border = border_cell
            if col_idx in [1, 5, 7]:
                cell.alignment = align_center
            else:
                cell.alignment = align_top_left
            if is_even:
                cell.fill = zebra_fill

    # =========================================================================
    # SHEET 2: Submateri (Topik)
    # =========================================================================
    ws2 = wb.create_sheet(title="2. Submateri (Topik)")
    ws2.views.sheetView[0].showGridLines = True
    
    ws2["A1"] = "FORMAT SUBMATERI & TOPIK PEMBELAJARAN"
    ws2["A1"].font = font_title
    ws2["A2"] = "Panduan: Setiap baris mewakili satu topik/submateri yang bernaung di bawah salah satu Modul Utama."
    ws2["A2"].font = font_italic
    
    headers_topic = [
        "Nomor Modul Target*", "Kode Topik*", "Urutan (Order Index)*", 
        "Judul Submateri*", "YouTube Video ID (Opsional)", "Isi Konten Ilmiah (HTML / Teks)*"
    ]
    
    for col_num, h in enumerate(headers_topic, 1):
        cell = ws2.cell(row=4, column=col_num, value=h)
        cell.font = font_header
        cell.fill = header_fill_teal
        cell.alignment = align_center
        cell.border = border_cell
        ws2.row_dimensions[4].height = 28
        
    sample_topics = [
        (1, "1.1", 1, "Pengenalan Anatomi Organ Reproduksi Pria & Wanita", "v3F4QW_h32M", 
         "<p>Sistem reproduksi manusia dirancang dengan fungsi biologis yang saling melengkapi. Memahami organ internal dan eksternal secara anatomis adalah langkah pertama dalam menjaga kesehatan organ vital secara mandiri.</p>"),
        (1, "1.2", 2, "Pubertas & Mekanisme Hormonal (Mimpi Basah & Siklus Haid)", "e7zP-b-YnS0", 
         "<p>Pubertas ditandai oleh kematangan aksis Hipotalamus-Hipofisis-Gonad yang memicu pelepasan hormon gonadotropin (GnRH, LH, FSH).</p>"),
        (1, "1.3", 3, "Variasi Karakteristik Darah Menstruasi", "", 
         "<p>Warna dan tekstur darah haid merupakan indikator penting keseimbangan hormonal serta sirkulasi darah di endometrium.</p>"),
        (1, "1.4", 4, "Higienitas Genitalia & Penanganan Dismenore (Skala NRS)", "", 
         "<p>Membasuh organ kewanitaan selalu dari arah depan ke belakang. Hindari penggunaan sabun antiseptik pewangi secara rutin.</p>"),
        (2, "2.1", 1, "Batasan Fisik & Emosional (Healthy Dating vs Relasi Toksik)", "", 
         "<p>Relasi yang sehat berlandaskan pada Mutual Respect, Trust, dan Consent. Kenali tanda bahaya (red flags) dalam pacaran.</p>"),
        (2, "2.2", 2, "Bedah Kasus Lingkungan Mahasiswa (Living Together & Kos-kosan)", "", 
         "<p>Kondisi indekos yang minim pengawasan sosial rentan terjadinya kohabitasi tanpa ikatan hukum resmi.</p>"),
        (2, "2.3", 3, "Keterampilan Asertif Remaja (Teknik Berkata 'TIDAK')", "", 
         "<p>Keterampilan asertif adalah kemampuan mengekspresikan batasan pribadi secara jujur, lugas, dan tenang tanpa agresif.</p>"),
        (5, "5.1", 1, "[Contoh Submateri Baru: Panduan Isi Piringku]", "dQw4w9WgXcQ", 
         "<p>Pemenuhan zat gizi makro dan mikro bagi remaja sangat krusial dalam metabolisme pembentukan sel darah merah.</p>")
    ]
    
    for row_idx, top in enumerate(sample_topics, 5):
        ws2.row_dimensions[row_idx].height = 40
        is_even = (row_idx % 2 == 0)
        for col_idx, val in enumerate(top, 1):
            cell = ws2.cell(row=row_idx, column=col_idx, value=val)
            cell.font = font_normal
            cell.border = border_cell
            if col_idx in [1, 2, 3, 5]:
                cell.alignment = align_center
            else:
                cell.alignment = align_top_left
            if is_even:
                cell.fill = zebra_fill

    # =========================================================================
    # SHEET 3: Bank Soal Evaluasi
    # =========================================================================
    ws3 = wb.create_sheet(title="3. Bank Soal Kuesioner")
    ws3.views.sheetView[0].showGridLines = True
    
    ws3["A1"] = "FORMAT BANK SOAL KUESIONER (PRE-TEST & POST-TEST)"
    ws3["A1"].font = font_title
    ws3["A2"] = "Panduan: Soal digunakan untuk evaluasi pemahaman modul dan pengukuran efektivitas edukasi (skor N-Gain)."
    ws3["A2"].font = font_italic
    
    headers_soal = [
        "Nomor Modul Target*", "Nomor Soal*", "Teks Pertanyaan Kuesioner*", 
        "Pilihan Jawaban A*", "Pilihan Jawaban B*", "Pilihan Jawaban C*", "Pilihan Jawaban D*", 
        "Kunci Jawaban (A/B/C/D)*", "Penjelasan Ilmiah / Pembahasan*"
    ]
    
    for col_num, h in enumerate(headers_soal, 1):
        cell = ws3.cell(row=4, column=col_num, value=h)
        cell.font = font_header
        cell.fill = header_fill_teal
        cell.alignment = align_center
        cell.border = border_cell
        ws3.row_dimensions[4].height = 28
        
    sample_soal = [
        (1, 1, "Arah yang benar saat membersihkan area organ kewanitaan setelah buang air adalah...",
         "Dari belakang (anus) ke arah depan (vagina)",
         "Dari depan (vagina) ke arah belakang (anus)",
         "Bebas dari arah mana saja yang penting menggunakan sabun wangi",
         "Cukup disiram tanpa perlu dibersihkan dengan tangan",
         "B",
         "Membasuh dari depan ke belakang mencegah perpindahan bakteri usus (seperti E. Coli) dari anus ke saluran kemih dan vagina."),
        (1, 2, "Darah menstruasi berwarna cokelat gelap atau kehitaman pada hari-hari terakhir siklus menunjukkan...",
         "Tanda infeksi menular seksual berbahaya",
         "Darah yang mengalami oksidasi normal saat keluar perlahan",
         "Kerusakan permanen pada dinding rahim",
         "Kekurangan vitamin dan mineral akut",
         "B",
         "Warna cokelat atau kehitaman adalah darah sisa yang telah teroksidasi oleh oksigen karena membutuhkan waktu lebih lama untuk keluar dari rahim, kondisi ini fisiologis/normal."),
        (2, 3, "Tindakan memanipulasi pasangan hingga membuat pasangan meragukan ingatan, kewarasan, atau perasaannya sendiri disebut...",
         "Love Language",
         "Assertive Communication",
         "Gaslighting (Manipulasi Emosional)",
         "Time Management",
         "C",
         "Gaslighting adalah bentuk pelecehan emosional di mana pelaku membuat korban mempertanyakan realitas atau perasaan mereka sendiri."),
        (3, 4, "Berikut ini yang merupakan komplikasi medis fatal akibat praktik aborsi tidak aman (unsafe abortion) adalah...",
         "Sepsis (infeksi berat sistemik) dan perforasi dinding rahim",
         "Peningkatan daya tahan tubuh terhadap bakteri",
         "Regenerasi sel sel telur lebih cepat",
         "Penurunan risiko anemia di masa depan",
         "A",
         "Aborsi ilegal tanpa prosedur medis steril sering menyebabkan robekan rahim (perforasi), perdarahan masif, infeksi darah mematikan (sepsis), dan infertilitas."),
        (4, 5, "Menurut BKKBN, batas usia minimal yang direkomendasikan untuk menikah agar organ reproduksi dan psikososial matang adalah...",
         "17 tahun wanita & 19 tahun pria",
         "21 tahun wanita & 25 tahun pria",
         "18 tahun wanita & 20 tahun pria",
         "15 tahun wanita & 18 tahun pria",
         "B",
         "BKKBN menganjurkan usia minimal 21 tahun untuk perempuan dan 25 tahun untuk laki-laki guna memastikan kesiapan biologis, kognitif, dan kemandirian finansial.")
    ]
    
    for row_idx, q in enumerate(sample_soal, 5):
        ws3.row_dimensions[row_idx].height = 42
        is_even = (row_idx % 2 == 0)
        for col_idx, val in enumerate(q, 1):
            cell = ws3.cell(row=row_idx, column=col_idx, value=val)
            cell.font = font_normal
            cell.border = border_cell
            if col_idx in [1, 2, 8]:
                cell.alignment = align_center
            else:
                cell.alignment = align_top_left
            if is_even:
                cell.fill = zebra_fill

    # =========================================================================
    # SHEET 4: Petunjuk Pengisian
    # =========================================================================
    ws4 = wb.create_sheet(title="4. Petunjuk & Ketentuan")
    ws4.views.sheetView[0].showGridLines = True
    
    ws4["A1"] = "PETUNJUK PENGISIAN TEMPLATE FORMAT MODUL SIKERA"
    ws4["A1"].font = font_title
    ws4["A2"] = "Baca aturan berikut sebelum mengisi atau mengunggah data modul baru ke platform SIKERA."
    ws4["A2"].font = font_italic
    
    guide_headers = ["Bagian", "Kolom", "Wajib / Opsional", "Format / Nilai yang Diterima", "Keterangan & Tips Praktis"]
    for col_num, h in enumerate(guide_headers, 1):
        cell = ws4.cell(row=4, column=col_num, value=h)
        cell.font = font_header
        cell.fill = header_fill_slate
        cell.alignment = align_center
        cell.border = border_cell
        ws4.row_dimensions[4].height = 26
        
    guidelines_data = [
        ("1. Modul Utama", "Nomor Modul", "Wajib", "Angka bulat (1, 2, 3, ...)", "Harus unik. Menentukan urutan modul di aplikasi mobile/web."),
        ("1. Modul Utama", "Judul Modul", "Wajib", "Teks singkat (maks 255 karakter)", "Nama materi utama modul edukasi."),
        ("1. Modul Utama", "Estimasi Waktu", "Wajib", "Teks (Contoh: '15 Menit')", "Perkiraan durasi mahasiswa menyelesaikan materi."),
        ("1. Modul Utama", "Path Banner", "Opsional", "Teks path file gambar (SVG/PNG)", "Jika dikosongkan, sistem akan otomatis memilih banner visual bawaan."),
        ("2. Submateri", "Kode Topik", "Wajib", "Teks (Contoh: '1.1', '1.2')", "Format nomor_modul.nomor_topik untuk penomoran bab."),
        ("2. Submateri", "YouTube Video ID", "Opsional", "11 Karakter ID Video (bukan URL)", "Contoh: jika URL https://youtu.be/v3F4QW_h32M maka isikan 'v3F4QW_h32M'."),
        ("2. Submateri", "Isi Konten", "Wajib", "Teks biasa atau tag HTML valid", "Dapat menggunakan tag <p>, <ul>, <li>, <strong>, <em>, <h4> untuk format rapi."),
        ("3. Bank Soal", "Nomor Modul Target", "Wajib", "Angka bulat (1, 2, 3, ...)", "Harus sesuai dengan nomor modul pada lembar '1. Modul Utama'."),
        ("3. Bank Soal", "Kunci Jawaban", "Wajib", "Satu huruf: A, B, C, atau D", "Kunci jawaban benar untuk scoring otomatis pre-test dan post-test."),
        ("3. Bank Soal", "Pembahasan", "Wajib", "Teks penjelasan ilmiah", "Tampil bagi mahasiswa setelah menyelesaikan evaluasi post-test.")
    ]
    
    for row_idx, row_data in enumerate(guidelines_data, 5):
        ws4.row_dimensions[row_idx].height = 26
        is_even = (row_idx % 2 == 0)
        for col_idx, val in enumerate(row_data, 1):
            cell = ws4.cell(row=row_idx, column=col_idx, value=val)
            cell.font = font_normal
            cell.border = border_cell
            if col_idx in [1, 3]:
                cell.alignment = align_center
            else:
                cell.alignment = align_left
            if is_even:
                cell.fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

    # Adjust column widths automatically
    for sheet in [ws1, ws2, ws3, ws4]:
        for col in sheet.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or '')
                if cell.row in [1, 2]:
                    continue
                # take first line if multiple lines
                line = val.split('\n')[0]
                if len(line) > max_len:
                    max_len = len(line)
            sheet.column_dimensions[col_letter].width = max(min(max_len + 4, 45), 14)
            
    # Custom specific column adjustments
    ws1.column_dimensions['A'].width = 16
    ws1.column_dimensions['B'].width = 32
    ws1.column_dimensions['C'].width = 35
    ws1.column_dimensions['D'].width = 45
    ws1.column_dimensions['E'].width = 18
    
    ws2.column_dimensions['A'].width = 20
    ws2.column_dimensions['B'].width = 14
    ws2.column_dimensions['C'].width = 18
    ws2.column_dimensions['D'].width = 35
    ws2.column_dimensions['E'].width = 24
    ws2.column_dimensions['F'].width = 50

    ws3.column_dimensions['A'].width = 20
    ws3.column_dimensions['B'].width = 14
    ws3.column_dimensions['C'].width = 45
    ws3.column_dimensions['D'].width = 28
    ws3.column_dimensions['E'].width = 28
    ws3.column_dimensions['F'].width = 28
    ws3.column_dimensions['G'].width = 28
    ws3.column_dimensions['H'].width = 20
    ws3.column_dimensions['I'].width = 45

    ws4.column_dimensions['A'].width = 18
    ws4.column_dimensions['B'].width = 24
    ws4.column_dimensions['C'].width = 16
    ws4.column_dimensions['D'].width = 32
    ws4.column_dimensions['E'].width = 50

    wb.save(output_file)
    print(f"Successfully generated: {output_file}")

if __name__ == '__main__':
    build_excel()
